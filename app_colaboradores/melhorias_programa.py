# melhorias_programa.py

try:
    from PyQt5 import QtWidgets, QtGui
except ImportError:
    QtWidgets = None
    QtGui = None
import importlib
import logging
import os
import sys
import tempfile
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path
import pandas as pd

from app_config import APP_DIR, RESOURCE_DIR

ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from agricola_shared.report_security import neutralize_dataframe


def aplicar_melhorias_globais(app):
    """
    Aplica melhorias globais ao QApplication: fonte padrão e estilo Fusion.
    """
    if QtWidgets is None or QtGui is None:
        raise RuntimeError("PyQt5 nao esta instalado neste ambiente.")
    QtWidgets.QApplication.setFont(QtGui.QFont("Segoe UI", 10))
    app.setStyle("Fusion")


def configurar_logger():
    """
    Configura o logger global para registrar logs no terminal e em arquivo.
    """
    logging.basicConfig(
        level=logging.INFO,
        format='%(asctime)s - %(levelname)s - %(message)s',
        handlers=[
            logging.FileHandler("log_sistema_colaboradores.log", encoding="utf-8"),
            logging.StreamHandler()
        ]
    )
    logging.info("Logger iniciado com sucesso.")


def verificar_dependencias_obrigatorias():
    """
    Verifica se bibliotecas essenciais estão instaladas.
    """
    dependencias_obrigatorias = {
        "fpdf": "Geração de PDFs",
        "matplotlib": "Dashboard",
        "qtawesome": "Ícones da interface",
    }
    dependencias_opcionais = {
        "xlsxwriter": "Exportação Excel com formatação avançada",
    }

    faltantes = []

    for modulo, uso in dependencias_obrigatorias.items():
        try:
            importlib.import_module(modulo)
        except ImportError:
            faltantes.append(f"{modulo} ({uso})")

    for modulo, uso in dependencias_opcionais.items():
        try:
            importlib.import_module(modulo)
        except ImportError:
            logging.info("Dependência opcional ausente: %s (%s).", modulo, uso)

    if faltantes:
        logging.error("Dependências obrigatórias ausentes: %s", ", ".join(faltantes))
    else:
        logging.info("Dependências obrigatórias carregadas com sucesso.")


def _resolver_caminho_tema(tema: str) -> str | None:
    nome_arquivo = "dark_theme.qss" if tema == "escuro" else "style.qss"
    base_dir = os.path.dirname(os.path.abspath(__file__))
    candidatos = [
        os.path.join(str(APP_DIR), nome_arquivo),
        os.path.join(str(APP_DIR), "assets", nome_arquivo),
        os.path.join(str(RESOURCE_DIR), nome_arquivo),
        os.path.join(str(RESOURCE_DIR), "assets", nome_arquivo),
        os.path.join(base_dir, nome_arquivo),
        os.path.join(base_dir, "assets", nome_arquivo),
        nome_arquivo,
        os.path.join("assets", nome_arquivo),
    ]
    for caminho in candidatos:
        if os.path.exists(caminho):
            return caminho
    return None


def carregar_qss_personalizado(janela, tema: str = "claro"):
    """
    Carrega o tema visual salvo pelo usuário.
    """
    qss_path = _resolver_caminho_tema(tema)
    if qss_path:
        with open(qss_path, "r", encoding="utf-8") as f:
            janela.setStyleSheet(f.read())
        logging.info(f"Tema '{tema}' aplicado com sucesso.")
    else:
        logging.warning("Arquivo de tema '%s' não encontrado.", tema)


def corrigir_nome_arquivo_excel(nome: str, nomes_usados: set[str] | None = None) -> str:
    """
    Corrige nomes inválidos para abas do Excel e, opcionalmente, garante
    unicidade sem diferenciar maiúsculas/minúsculas.
    """
    invalidos = ['/', '\\', '*', '?', ':', '[', ']']
    nome = str(nome or "").strip().strip("'")
    for c in invalidos:
        nome = nome.replace(c, '-')
    nome = nome.strip().strip("'") or "Planilha"
    nome = nome[:31]

    if nomes_usados is None:
        return nome

    usados_normalizados = {str(item).casefold() for item in nomes_usados}
    candidato = nome
    indice = 2
    while candidato.casefold() in usados_normalizados:
        sufixo = f" ({indice})"
        candidato = f"{nome[:31 - len(sufixo)]}{sufixo}"
        indice += 1
    nomes_usados.add(candidato)
    return candidato


@contextmanager
def caminho_publicacao_atomica(caminho: str | os.PathLike):
    """Entrega um arquivo temporário irmão e só publica ao final com os.replace."""
    destino = Path(caminho)
    destino.parent.mkdir(parents=True, exist_ok=True)
    descritor, temporario = tempfile.mkstemp(
        prefix=f".{destino.stem}.",
        suffix=f".tmp{destino.suffix}",
        dir=str(destino.parent),
    )
    os.close(descritor)
    temporario_path = Path(temporario)
    try:
        yield str(temporario_path)
        if not temporario_path.is_file() or temporario_path.stat().st_size <= 0:
            raise RuntimeError(f"Arquivo temporário não foi gerado corretamente: {temporario_path}")
        os.replace(temporario_path, destino)
    finally:
        try:
            temporario_path.unlink(missing_ok=True)
        except OSError:
            logging.warning("Não foi possível remover temporário de relatório: %s", temporario_path)


def _coluna_excel_tem_datas(serie: pd.Series) -> bool:
    if pd.api.types.is_datetime64_any_dtype(serie.dtype):
        return True
    valores = serie.dropna()
    return bool(not valores.empty and valores.map(lambda valor: isinstance(valor, (date, datetime, pd.Timestamp))).all())


def salvar_df_excel_com_total(
    df: pd.DataFrame,
    caminho: str,
    nome_aba: str = "Relatório",
    *,
    incluir_totais: bool = True,
):
    """
    Salva um DataFrame em Excel com linha de totais, se houver colunas numéricas.
    """
    try:
        import xlsxwriter  # noqa: F401
        engine = "xlsxwriter"
    except ImportError:
        engine = "openpyxl"

    df = neutralize_dataframe(df.copy())
    nome_aba = corrigir_nome_arquivo_excel(nome_aba)
    writer_kwargs = {"engine_kwargs": {"options": {"strings_to_formulas": False}}} if engine == "xlsxwriter" else {}
    with caminho_publicacao_atomica(caminho) as caminho_temporario:
        with pd.ExcelWriter(
            caminho_temporario,
            engine=engine,
            date_format="dd/mm/yyyy",
            datetime_format="dd/mm/yyyy",
            **writer_kwargs,
        ) as writer:
            df.to_excel(writer, sheet_name=nome_aba, index=False, startrow=1, header=False)
            worksheet = writer.sheets[nome_aba]

            if engine == "xlsxwriter":
                workbook = writer.book
                header_format = workbook.add_format({
                    'bold': True, 'font_color': '#FFFFFF', 'bg_color': '#202467',
                    'border': 1, 'border_color': '#D9E2DC', 'text_wrap': True,
                    'valign': 'vcenter',
                })
                date_format = workbook.add_format({'num_format': 'dd/mm/yyyy', 'valign': 'top'})
                wrap_format = workbook.add_format({'text_wrap': True, 'valign': 'top'})
                for col_num, value in enumerate(df.columns.values):
                    worksheet.write(0, col_num, value, header_format)
                worksheet.freeze_panes(1, 0)
                worksheet.hide_gridlines(2)
                if len(df.columns):
                    worksheet.autofilter(0, 0, max(len(df), 0), len(df.columns) - 1)
                worksheet.set_row(0, 30)
            else:
                from openpyxl.styles import Alignment, Font, PatternFill

                header_fill = PatternFill(fill_type="solid", fgColor="202467")
                for col_num, value in enumerate(df.columns.values, start=1):
                    cell = worksheet.cell(row=1, column=col_num, value=value)
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = header_fill
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                worksheet.freeze_panes = "A2"
                worksheet.sheet_view.showGridLines = False
                if len(df.columns):
                    from openpyxl.utils import get_column_letter
                    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{max(1, len(df) + 1)}"
                worksheet.row_dimensions[1].height = 30

            numeric_cols = df.select_dtypes(include='number').columns
            if incluir_totais and not numeric_cols.empty:
                total_row = df[numeric_cols].sum()
                total_row_dict = {col: total_row[col] for col in numeric_cols}
                pd.DataFrame([total_row_dict]).to_excel(
                    writer,
                    sheet_name=nome_aba,
                    index=False,
                    startrow=len(df) + 2,
                    header=False,
                )

            for i, col in enumerate(df.columns):
                valores_texto = (
                    df[col].map(lambda valor: "" if pd.isna(valor) else str(valor))
                    if not df.empty else pd.Series(dtype=str)
                )
                maior_valor = valores_texto.map(len).max() if not valores_texto.empty else 0
                # XlsxWriter acrescenta cerca de 0,71 à largura lida pelo Excel;
                # limitar a 40 mantém o arquivo final confortavelmente abaixo de 42.
                width = min(max(maior_valor, len(str(col))) + 2, 40)
                coluna_data = _coluna_excel_tem_datas(df[col])
                if coluna_data:
                    width = min(max(width, 13), 18)
                if engine == "xlsxwriter":
                    worksheet.set_column(i, i, width, date_format if coluna_data else wrap_format)
                else:
                    from openpyxl.styles import Alignment
                    from openpyxl.utils import get_column_letter

                    worksheet.column_dimensions[get_column_letter(i + 1)].width = width
                    for cell in worksheet.iter_cols(
                        min_col=i + 1,
                        max_col=i + 1,
                        min_row=2,
                        max_row=max(2, worksheet.max_row),
                    ):
                        for item in cell:
                            item.alignment = Alignment(wrap_text=True, vertical="top")
                            if coluna_data and item.value not in (None, ""):
                                item.number_format = "dd/mm/yyyy"

        logging.info(f"Relatório Excel salvo em: {caminho}")
