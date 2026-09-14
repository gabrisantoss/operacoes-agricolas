import atexit
from datetime import datetime
import importlib
import os
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest import mock

import pandas as pd
from fpdf import FPDF
from openpyxl import load_workbook
from PyPDF2 import PdfReader


def _configure_isolated_runtime() -> None:
    existing = os.environ.get("APP_COLAB_TEST_RUNTIME")
    if existing:
        runtime = Path(existing)
    else:
        runtime = Path(tempfile.mkdtemp(prefix="app-colaboradores-ui-report-tests-"))
        os.environ["APP_COLAB_TEST_RUNTIME"] = str(runtime)
        atexit.register(shutil.rmtree, runtime, ignore_errors=True)
    os.environ["APP_COLAB_CONFIG_FILE"] = str(runtime / "missing-config.json")
    os.environ["APP_COLAB_DB_ENGINE"] = "sqlite"
    os.environ["APP_COLAB_SQLITE_PATH"] = str(runtime / "fixture.db")
    os.environ["APP_COLAB_STORAGE_ROOT"] = str(runtime / "storage")
    os.environ["APP_COLAB_BACKUP_DIR"] = str(runtime / "backups")
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


_configure_isolated_runtime()

from PyQt5 import QtWidgets  # noqa: E402

profile_window_module = importlib.import_module("ui.ProfileWindow")
from ui.ProfileWindow import _draw_profile_pdf_row  # noqa: E402
from ui.TabIDO import _prepare_ido_export_dataframe, _write_ido_excel  # noqa: E402
from ui.TabRelatorios import PDFRelatorio, TabRelatorios  # noqa: E402


class TabRelatoriosPdfTests(unittest.TestCase):
    def test_multiline_grid_repeats_group_and_header_on_every_page(self):
        with tempfile.TemporaryDirectory() as temporary:
            destination = Path(temporary) / "relatorio_multilinha.pdf"
            pdf = PDFRelatorio(titulo="Relatório de QA")
            widths = [80, 30, 40, 30]
            headers = ["Nome", "Código", "Função", "Horário"]
            pdf.set_table_context(
                widths,
                headers,
                section_title="Frente QA - Turno A",
            )
            pdf.add_page()
            pdf.set_font("Arial", "", 9)

            for index in range(75):
                TabRelatorios._draw_multi_line_row(
                    None,
                    pdf,
                    widths,
                    [
                        f"Colaborador de validação com nome muito extenso número {index:03d}",
                        f"C{index:03d}",
                        "Operador de Máquinas Agrícolas e Equipamentos",
                        "07:00 às 15:20",
                    ],
                )

            pdf.output(str(destination))
            pages = PdfReader(str(destination)).pages

            self.assertGreater(len(pages), 1)
            for page in pages:
                text = page.extract_text()
                self.assertIn("Frente QA - Turno A", text)
                self.assertIn("Nome", text)
                self.assertIn("Código", text)
                self.assertIn("Função", text)
                self.assertIn("Horário", text)

    def test_profile_row_wraps_and_moves_whole_row_to_next_page(self):
        with tempfile.TemporaryDirectory() as temporary:
            destination = Path(temporary) / "ficha_multilinha.pdf"
            pdf = FPDF()
            pdf.set_auto_page_break(auto=True, margin=20)
            pdf.add_page()
            pdf.set_y(pdf.page_break_trigger - 3)
            long_value = "INÍCIO " + ("observação extensa " * 28) + " FIM"

            height = _draw_profile_pdf_row(
                pdf,
                "Observações",
                long_value,
                continuation_section="Dados Profissionais",
            )
            pdf.output(str(destination))

            self.assertGreater(height, 8)
            self.assertEqual(2, pdf.page_no())
            text = PdfReader(str(destination)).pages[-1].extract_text()
            self.assertIn("Dados Profissionais", text)
            self.assertIn("INÍCIO", text)
            self.assertIn("FIM", text)

    def test_profile_pdf_keeps_section_with_first_row_and_repeats_it_after_break(self):
        with tempfile.TemporaryDirectory() as temporary:
            destination = Path(temporary) / "ficha_multpagina.pdf"
            app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
            self.addCleanup(app.processEvents)
            with mock.patch.object(profile_window_module.ProfileWindow, "_setup_ui"), mock.patch.object(
                profile_window_module.ProfileWindow,
                "carregar_dados",
            ):
                profile = profile_window_module.ProfileWindow("QA-STRESS")

            long_value = (
                "Texto sintético extenso para validar quebra de linha, grade e paginação segura. "
                * 12
            )
            profile.colab_data = {
                "nome": f"Colaborador Sintético {long_value}",
                "cpf": "000.000.000-00",
                "rg": "00.000.000-0",
                "nascimento": "01/01/1990",
                "municipio": long_value,
                "telefone": "(16) 99999-0000",
                "codigo_colaborador": "QA-STRESS",
                "funcao": long_value,
                "data_admissao": "01/01/2020",
                "local_trabalho": long_value,
                "gestor_responsavel": long_value,
                "salario": 1234.56,
                "registro_cnh": "00000000000",
                "categoria_cnh": "D",
                "validade_cnh": "15/09/2026",
                "status_cnh_acompanhamento": "EM_ANDAMENTO",
            }

            with mock.patch.object(
                QtWidgets.QFileDialog,
                "getSaveFileName",
                return_value=(str(destination), ""),
            ), mock.patch.object(
                QtWidgets.QMessageBox,
                "information",
                return_value=None,
            ), mock.patch.object(
                QtWidgets.QMessageBox,
                "critical",
                return_value=None,
            ) as critical, mock.patch.object(profile_window_module.sys, "platform", "linux"):
                profile.gerar_ficha_pdf()

            critical.assert_not_called()
            texts = [page.extract_text() or "" for page in PdfReader(str(destination)).pages]
            self.assertGreater(len(texts), 1)
            matricula_page = next(text for text in texts if "Matrícula (Código)" in text)
            gestor_page = next(text for text in texts if "Gestor Responsável" in text)
            self.assertIn("Dados Profissionais", matricula_page)
            self.assertIn("Dados Profissionais", gestor_page)


class IdoWorkbookTests(unittest.TestCase):
    def _fixture(self) -> pd.DataFrame:
        return pd.DataFrame(
            [
                {
                    "Nome": "=2+2",
                    "Código": "+123",
                    "Matrícula": "@usuario",
                    "Seção": "Operação agrícola com descrição longa para largura",
                    "Total de Informes": "7",
                    "Período": "08/2026",
                }
            ]
        )

    def test_ido_export_has_safe_strings_real_types_filters_and_bounded_widths(self):
        with tempfile.TemporaryDirectory() as temporary:
            destination = Path(temporary) / "ido.xlsx"
            dataframe = _prepare_ido_export_dataframe(self._fixture())
            _write_ido_excel(dataframe, destination)

            workbook = load_workbook(destination, data_only=False)
            worksheet = workbook["IDO"]
            headers = {cell.value: cell.column for cell in worksheet[1]}

            self.assertEqual("A2", str(worksheet.freeze_panes))
            self.assertEqual("A1:F2", worksheet.auto_filter.ref)
            self.assertEqual(7, worksheet.cell(2, headers["Total de Informes"]).value)
            self.assertEqual("n", worksheet.cell(2, headers["Total de Informes"]).data_type)

            period_cell = worksheet.cell(2, headers["Período"])
            self.assertIsInstance(period_cell.value, datetime)
            self.assertEqual("mm/yyyy", period_cell.number_format)

            for header in ("Nome", "Código", "Matrícula"):
                cell = worksheet.cell(2, headers[header])
                self.assertEqual("s", cell.data_type)
                self.assertTrue(str(cell.value).startswith("'"))

            self.assertFalse(
                any(
                    cell.data_type == "f"
                    for row in worksheet.iter_rows()
                    for cell in row
                )
            )
            self.assertTrue(
                all(
                    dimension.width is None or dimension.width <= 42
                    for dimension in worksheet.column_dimensions.values()
                )
            )
            self.assertEqual([], workbook._external_links)
            workbook.close()

    def test_ido_export_failure_preserves_previous_file_and_removes_staging_file(self):
        with tempfile.TemporaryDirectory() as temporary:
            destination = Path(temporary) / "ido.xlsx"
            destination.write_bytes(b"PREVIOUS")
            dataframe = _prepare_ido_export_dataframe(self._fixture())

            with mock.patch.object(
                pd.DataFrame,
                "to_excel",
                side_effect=RuntimeError("falha controlada"),
            ):
                with self.assertRaisesRegex(RuntimeError, "falha controlada"):
                    _write_ido_excel(dataframe, destination)

            self.assertEqual(b"PREVIOUS", destination.read_bytes())
            self.assertEqual([], list(destination.parent.glob(".ido.*.tmp.xlsx")))


if __name__ == "__main__":
    unittest.main()
