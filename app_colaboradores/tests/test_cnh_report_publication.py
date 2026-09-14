import atexit
from datetime import datetime
import os
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest import mock

from openpyxl import load_workbook
import pandas as pd


def _configure_isolated_runtime() -> None:
    existing = os.environ.get("APP_COLAB_TEST_RUNTIME")
    if existing:
        runtime = Path(existing)
    else:
        runtime = Path(tempfile.mkdtemp(prefix="app-colaboradores-report-tests-"))
        os.environ["APP_COLAB_TEST_RUNTIME"] = str(runtime)
        atexit.register(shutil.rmtree, runtime, ignore_errors=True)
    os.environ["APP_COLAB_CONFIG_FILE"] = str(runtime / "missing-config.json")
    os.environ["APP_COLAB_DB_ENGINE"] = "sqlite"
    os.environ["APP_COLAB_SQLITE_PATH"] = str(runtime / "fixture.db")
    os.environ["APP_COLAB_STORAGE_ROOT"] = str(runtime / "storage")
    os.environ["APP_COLAB_BACKUP_DIR"] = str(runtime / "backups")


_configure_isolated_runtime()

from core import cnh_management, cnh_reports  # noqa: E402
from melhorias_programa import salvar_df_excel_com_total  # noqa: E402


def _row(frente: str, codigo: str = "1") -> dict:
    return {
        "prioridade_label": "Alta",
        "prioridade_ordem": 1,
        "codigo_colaborador": codigo,
        "nome": f"Colaborador {codigo}",
        "cidade_exibicao": "Sertãozinho",
        "frente_exibicao": frente,
        "gestor_exibicao": "Gestor",
        "funcao_exibicao": "Operador",
        "turno_safra": "A",
        "horario": "07:00-15:00",
        "categoria_cnh": "D",
        "validade_cnh": "2026-09-15",
        "validade_cnh_formatada": "15/09/2026",
        "status_tecnico": "VENCIDA",
        "status_tecnico_label": "Vencida",
        "dias_para_vencer": -5,
        "status_acompanhamento_label": "Em andamento",
        "ultimo_contato_cnh": "2026-08-20",
        "responsavel_ultimo_contato_cnh": "RH",
        "data_prevista_regularizacao_cnh": "2026-09-10",
        "observacao_cnh": "Texto de acompanhamento",
        "inconsistencias_texto": "",
    }


def _pdf_row() -> cnh_reports.CNHReportRow:
    return cnh_reports.CNHReportRow(
        codigo_colaborador="1",
        nome="Colaborador",
        frente_safra="FRENTE 1",
        turno_safra="A",
        funcao_safra="Operador",
        validade_cnh="15/09/2026",
        dias_para_vencer=-5,
    )


class CnhReportPublicationTests(unittest.TestCase):
    def test_package_removes_stale_empty_categories(self):
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary)
            names = {
                "vencidas": "cnhs_vencidas_por_frente.pdf",
                "a_vencer": "cnhs_a_vencer_30_dias_por_frente.pdf",
                "validas": "cnhs_validas_por_frente.pdf",
            }
            for name in names.values():
                (output / name).write_bytes(b"OLD")
            old_alert_window = output / "cnhs_a_vencer_15_dias_por_frente.pdf"
            old_alert_window.write_bytes(b"OLD")

            fixture = {
                "categorias": {"vencidas": [_pdf_row()], "a_vencer": [], "validas": []},
                "resumo": {"total_lidos": 1, "sem_validade": 0, "datas_invalidas": 0, "sem_escala_ignorados": 0},
            }
            with mock.patch.object(cnh_reports, "carregar_registros_cnh", return_value=fixture):
                result = cnh_reports.gerar_relatorios_cnh("unused.db", str(output))

            self.assertTrue((output / names["vencidas"]).read_bytes().startswith(b"%PDF"))
            self.assertFalse((output / names["a_vencer"]).exists())
            self.assertFalse(old_alert_window.exists())
            self.assertFalse((output / names["validas"]).exists())
            self.assertIsNone(result["arquivos"]["a_vencer"])
            self.assertIsNone(result["arquivos"]["validas"])

    def test_pdf_failure_preserves_published_file_and_cleans_temp(self):
        with tempfile.TemporaryDirectory() as temporary:
            destination = Path(temporary) / "cnh.pdf"
            destination.write_bytes(b"PREVIOUS")
            with mock.patch.object(cnh_reports.PDFRelatorioCNH, "output", side_effect=RuntimeError("falha controlada")):
                with self.assertRaisesRegex(RuntimeError, "falha controlada"):
                    cnh_reports.gerar_pdf_categoria([_pdf_row()], str(destination), "Título", "Subtítulo")
            self.assertEqual(b"PREVIOUS", destination.read_bytes())
            self.assertEqual([], list(destination.parent.glob(".cnh.*.tmp.pdf")))

    def test_package_generation_failure_does_not_publish_partial_files(self):
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary)
            targets = [
                output / "cnhs_vencidas_por_frente.pdf",
                output / "cnhs_a_vencer_30_dias_por_frente.pdf",
                output / "cnhs_validas_por_frente.pdf",
            ]
            for target in targets:
                target.write_bytes(b"PREVIOUS")
            fixture = {
                "categorias": {"vencidas": [_pdf_row()], "a_vencer": [_pdf_row()], "validas": [_pdf_row()]},
                "resumo": {},
            }

            calls = 0

            def controlled_generation(_items, path, _title, _subtitle):
                nonlocal calls
                calls += 1
                if calls == 2:
                    raise RuntimeError("falha controlada")
                Path(path).write_bytes(b"STAGED")
                return path

            with mock.patch.object(cnh_reports, "carregar_registros_cnh", return_value=fixture), mock.patch.object(
                cnh_reports, "gerar_pdf_categoria", side_effect=controlled_generation
            ):
                with self.assertRaisesRegex(RuntimeError, "falha controlada"):
                    cnh_reports.gerar_relatorios_cnh("unused.db", str(output))

            self.assertTrue(all(target.read_bytes() == b"PREVIOUS" for target in targets))


class CnhWorkbookTests(unittest.TestCase):
    def test_generic_export_keeps_numbers_and_neutralizes_formula_like_text(self):
        dataframe = pd.DataFrame(
            [{"Nome": "=HYPERLINK(\"https://invalid.example\",\"x\")", "Quantidade": 7}]
        )
        with tempfile.TemporaryDirectory() as temporary:
            destination = Path(temporary) / "safe.xlsx"
            salvar_df_excel_com_total(dataframe, str(destination), "Painel")
            workbook = load_workbook(destination, data_only=False)
            sheet = workbook["Painel"]
            self.assertEqual("'=HYPERLINK(\"https://invalid.example\",\"x\")", sheet["A2"].value)
            self.assertEqual("s", sheet["A2"].data_type)
            self.assertEqual(7, sheet["B2"].value)
            self.assertEqual("n", sheet["B2"].data_type)
            self.assertTrue(all(dimension.width <= 42 for dimension in sheet.column_dimensions.values()))
            self.assertEqual([], workbook._external_links)
            workbook.close()

    def test_front_sheet_names_are_valid_unique_and_case_insensitive(self):
        linhas = [
            _row("A/B", "1"),
            _row("A:B", "2"),
            _row("'FRENTE 1'", "3"),
            _row("geral", "4"),
        ]
        with tempfile.TemporaryDirectory() as temporary:
            destination = Path(temporary) / "cobranca.xlsx"
            cnh_management._salvar_cobranca_por_frente_excel(linhas, str(destination))
            workbook = load_workbook(destination, read_only=False, data_only=False)
            names = workbook.sheetnames
            workbook.close()

        self.assertEqual(len(names), len({name.casefold() for name in names}))
        self.assertTrue(all(len(name) <= 31 for name in names))
        self.assertTrue(all(not any(char in name for char in "[]:*?/\\") for name in names))
        self.assertTrue(all(not (name.startswith("'") or name.endswith("'")) for name in names))
        self.assertIn("A-B", names)
        self.assertIn("A-B (2)", names)
        self.assertIn("FRENTE 1", names)
        self.assertIn("geral (2)", names)

    def test_cnh_dates_are_real_cells_and_tables_are_usable(self):
        linhas = [_row("FRENTE 1")]
        with tempfile.TemporaryDirectory() as temporary:
            charge_path = Path(temporary) / "cobranca.xlsx"
            panel_path = Path(temporary) / "painel.xlsx"
            cnh_management._salvar_cobranca_por_frente_excel(linhas, str(charge_path))
            panel = cnh_management._montar_dataframe_operacional(linhas)
            salvar_df_excel_com_total(panel, str(panel_path), "Painel CNH")

            charge = load_workbook(charge_path, data_only=False)
            for sheet_name in ("Geral", "FRENTE 1"):
                sheet = charge[sheet_name]
                headers = {cell.value: cell.column for cell in sheet[5]}
                self.assertEqual("A6", str(sheet.freeze_panes))
                self.assertTrue(sheet.auto_filter.ref)
                for header in ("Validade CNH", "Último contato", "Próxima ação"):
                    cell = sheet.cell(6, headers[header])
                    self.assertIsInstance(cell.value, datetime)
                    self.assertEqual("dd/mm/yyyy", cell.number_format)
                self.assertTrue(all(dimension.width <= 42 for dimension in sheet.column_dimensions.values()))
            charge.close()

            panel_book = load_workbook(panel_path, data_only=False)
            sheet = panel_book["Painel CNH"]
            headers = {cell.value: cell.column for cell in sheet[1]}
            self.assertEqual("A2", str(sheet.freeze_panes))
            self.assertTrue(sheet.auto_filter.ref)
            for header in ("Validade CNH", "Último Contato", "Data Prevista"):
                self.assertIsInstance(sheet.cell(2, headers[header]).value, datetime)
                self.assertEqual("dd/mm/yyyy", sheet.cell(2, headers[header]).number_format)
            self.assertTrue(all(dimension.width <= 42 for dimension in sheet.column_dimensions.values()))
            panel_book.close()

    def test_xlsx_failure_preserves_previous_destination(self):
        with tempfile.TemporaryDirectory() as temporary:
            destination = Path(temporary) / "cobranca.xlsx"
            destination.write_bytes(b"PREVIOUS")
            with mock.patch.object(
                cnh_management,
                "_aplicar_formato_tabela_cobranca",
                side_effect=RuntimeError("falha controlada"),
            ):
                with self.assertRaisesRegex(RuntimeError, "falha controlada"):
                    cnh_management._salvar_cobranca_por_frente_excel([_row("FRENTE 1")], str(destination))
            self.assertEqual(b"PREVIOUS", destination.read_bytes())
            self.assertEqual([], list(destination.parent.glob(".cobranca.*.tmp.xlsx")))


if __name__ == "__main__":
    unittest.main()
