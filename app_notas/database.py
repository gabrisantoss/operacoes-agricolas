from __future__ import annotations

import json
import hashlib
import os
import shutil
import sqlite3
import subprocess
import sys
import threading
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Iterable
from urllib.parse import unquote, urlparse

import psycopg

from app_config import APP_ROOT, DATABASE_URL, DB_ENGINE, DB_PATH, EXCEL_ORIGEM, SQLITE_BUSY_TIMEOUT_MS
from app_logging import get_logger

ROOT_DIR = APP_ROOT.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from agricola_shared.db_compat import connect as connect_postgres_db
from agricola_shared.db_compat import connect_sqlite as connect_sqlite_db
from agricola_shared.db_compat import translate_sql

try:
    from colaboradores_reference import get_colaboradores_provider
except ModuleNotFoundError:
    get_colaboradores_provider = None

try:
    from openpyxl import load_workbook
except ModuleNotFoundError:
    load_workbook = None


LOGGER = get_logger(__name__)
DATABASE_ERROR_TYPES = (sqlite3.DatabaseError, psycopg.Error)
INTEGRITY_ERROR_TYPES = (sqlite3.IntegrityError, psycopg.IntegrityError)

# Ordenacao numerica de `codigo` (TEXT) portavel entre engines.
# SQLite: CAST(...) e leniente (texto nao numerico vira 0).
# PostgreSQL: CAST AS INTEGER lanca erro em texto nao numerico, entao extraimos
# apenas os digitos; codigos sem nenhum digito vao para o fim (NULLS LAST).
if DB_ENGINE in {"postgres", "postgresql"}:
    CODIGO_NUMERIC_ORDER = "NULLIF(regexp_replace(codigo, '[^0-9]', '', 'g'), '')::bigint NULLS LAST"
else:
    CODIGO_NUMERIC_ORDER = "CAST(codigo AS INTEGER)"

ABAS_EXCEL = {
    "motoristas": "nome",
    "fazendas": "fazendas",
    "variedades": "variedades",
}

NOTA_FIELDS = (
    "numero",
    "motorista_cod",
    "motorista_nome",
    "caminhao",
    "operador_cod",
    "operador_nome",
    "colhedora",
    "faz_muda_cod",
    "faz_muda_nome",
    "talhao",
    "faz_plantio_cod",
    "faz_plantio_nome",
    "variedade_id",
    "variedade_nome",
    "data_colheita",
    "data_plantio",
)

REFERENCE_TABLES = {
    "motoristas": {"id_col": "codigo", "label": "motorista"},
    "fazendas": {"id_col": "codigo", "label": "fazenda"},
    "variedades": {"id_col": "id", "label": "variedade"},
}

CORRUPTION_MARKERS = (
    "database disk image is malformed",
    "file is not a database",
    "database schema is malformed",
)


class DatabaseCorruptionError(sqlite3.DatabaseError):
    def __init__(self, path: str | Path, details: str):
        self.path = Path(path)
        self.details = str(details).strip() or "Falha desconhecida na integridade do banco."
        super().__init__(f"Banco corrompido em '{self.path}': {self.details}")


class CorrectionPreviewConflict(ValueError):
    """Raised when notes changed after the user reviewed a correction preview."""


def _is_corruption_message(message: str) -> bool:
    message_lc = str(message).lower()
    return any(marker in message_lc for marker in CORRUPTION_MARKERS)


def _sqlite_sidecar_paths(db_path: str | Path) -> tuple[Path, ...]:
    path = Path(db_path)
    return (
        path.with_name(f"{path.name}-wal"),
        path.with_name(f"{path.name}-shm"),
        path.with_name(f"{path.name}-journal")
    )


def sqlite_integrity_status(db_path: str | Path) -> str:
    path = Path(db_path)
    if not path.exists():
        raise FileNotFoundError(f"Banco de dados nao encontrado: {path}")

    conn = sqlite3.connect(path)
    try:
        row = conn.execute("PRAGMA quick_check;").fetchone()
    except sqlite3.DatabaseError as exc:
        return str(exc)
    finally:
        conn.close()

    if not row or row[0] is None:
        return "quick_check sem retorno"
    return str(row[0]).strip()


def ensure_sqlite_integrity(db_path: str | Path) -> None:
    status = sqlite_integrity_status(db_path)
    if status.lower() != "ok":
        raise DatabaseCorruptionError(db_path, status)


def _quarantine_sqlite_files(db_path: str | Path) -> dict[Path, Path]:
    path = Path(db_path)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    moved: dict[Path, Path] = {}

    for candidate in (path, *_sqlite_sidecar_paths(path)):
        if not candidate.exists():
            continue

        quarantined = candidate.with_name(f"{candidate.name}.corrompido_{timestamp}")
        sequence = 1
        while quarantined.exists():
            quarantined = candidate.with_name(
                f"{candidate.name}.corrompido_{timestamp}_{sequence}"
            )
            sequence += 1

        candidate.replace(quarantined)
        moved[candidate] = quarantined

    return moved


def _restore_quarantined_files(moved: dict[Path, Path]) -> None:
    for original, quarantined in reversed(list(moved.items())):
        if original.exists() or not quarantined.exists():
            continue
        quarantined.replace(original)


def _normalize_note_date(value, field_label: str) -> str | None:
    if value in (None, ""):
        return value

    if isinstance(value, datetime):
        parsed = value.date()
    elif isinstance(value, date):
        parsed = value
    else:
        try:
            parsed = date.fromisoformat(str(value).strip())
        except ValueError as exc:
            raise ValueError(f"{field_label} invalida: use o formato AAAA-MM-DD.") from exc

    today = date.today()
    if parsed > today:
        today_br = today.strftime("%d/%m/%Y")
        raise ValueError(
            f"{field_label} nao pode ser maior que a data atual do sistema ({today_br})."
        )

    return parsed.isoformat()


def restore_sqlite_backup(
    destination_path: str | Path,
    backup_path: str | Path,
) -> tuple[Path, Path | None]:
    destination = Path(destination_path)
    origem = Path(backup_path)
    if not origem.exists():
        raise FileNotFoundError(f"Backup nao encontrado: {origem}")

    ensure_sqlite_integrity(origem)
    destination.parent.mkdir(parents=True, exist_ok=True)

    moved = _quarantine_sqlite_files(destination)
    quarantined_main = moved.get(destination)

    try:
        backup_sqlite_file(origem, destination)
        ensure_sqlite_integrity(destination)
    except Exception:
        try:
            if destination.exists():
                destination.unlink()
        finally:
            _restore_quarantined_files(moved)
        raise

    return destination, quarantined_main


def backup_sqlite_file(source_path: str | Path, destination_path: str | Path) -> Path:
    source = Path(source_path)
    destination = Path(destination_path)

    if not source.exists():
        raise FileNotFoundError(f"Banco de dados nao encontrado: {source}")

    destination.parent.mkdir(parents=True, exist_ok=True)

    src = sqlite3.connect(source)
    dst = sqlite3.connect(destination)
    try:
        src.backup(dst)
        dst.commit()
    finally:
        dst.close()
        src.close()

    return destination


def _find_pg_tool(name: str) -> str:
    candidate = shutil.which(name)
    if candidate:
        return candidate
    program_files = Path(os.environ.get("ProgramFiles", r"C:\Program Files"))
    suffix = ".exe" if os.name == "nt" and not name.endswith(".exe") else ""
    for path in sorted(program_files.glob(f"PostgreSQL/*/bin/{name}{suffix}"), reverse=True):
        if path.exists():
            return str(path)
    raise RuntimeError(f"{name} nao encontrado para backup PostgreSQL.")


def run_pg_dump(database_url: str, destination: Path) -> None:
    from agricola_shared.demo_safety import assert_demo_database_target
    assert_demo_database_target(database_url)
    parsed = urlparse(database_url)
    env = os.environ.copy()
    env["PGPASSWORD"] = unquote(parsed.password or "")
    completed = subprocess.run(
        [
            _find_pg_tool("pg_dump"),
            "-h",
            parsed.hostname or "127.0.0.1",
            "-p",
            str(parsed.port or 5432),
            "-U",
            unquote(parsed.username or ""),
            "-d",
            parsed.path.lstrip("/"),
            "-F",
            "c",
            "-f",
            str(destination),
        ],
        env=env,
        text=True,
        capture_output=True,
        check=False
    )
    if completed.returncode != 0:
        raise RuntimeError(completed.stderr.strip() or completed.stdout.strip() or "pg_dump falhou")


class DB:
    def __init__(self, path: str | Path = DB_PATH, seed_from_excel: bool = True):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._lock = threading.RLock()
        need_seed = seed_from_excel and (
            not self.path.exists() or self.path.stat().st_size == 0
        )

        self.conn = self._open_connection()
        try:
            self._validate_integrity()
            self._criar_tabelas()
            self._migrar_schema()
            self._criar_indices()

            if need_seed:
                try:
                    self._seed_from_excel()
                except Exception:
                    LOGGER.exception("Falha ao importar dados iniciais do Excel")
        except Exception:
            self.close()
            raise

    def _open_connection(self) -> sqlite3.Connection:
        if DB_ENGINE in {"postgres", "postgresql"}:
            if not DATABASE_URL:
                raise RuntimeError("APP_NOTAS_DATABASE_URL nao configurada para PostgreSQL.")
            conn = connect_postgres_db(DATABASE_URL, row_factory=sqlite3.Row)
        else:
            try:
                conn = connect_sqlite_db(
                    self.path,
                    timeout=SQLITE_BUSY_TIMEOUT_MS / 1000,
                    row_factory=sqlite3.Row
                )
            except DATABASE_ERROR_TYPES as exc:
                # connect_sqlite ja roda PRAGMAs; um arquivo corrompido falha aqui
                # antes de _configure_connection. Embrulha como DatabaseCorruptionError.
                raise self._wrap_database_error(exc) from exc
        try:
            self._configure_connection(conn)
        except DATABASE_ERROR_TYPES as exc:
            conn.close()
            raise self._wrap_database_error(exc) from exc
        return conn

    def _validate_integrity(self) -> None:
        with self._lock:
            try:
                row = self.conn.execute("PRAGMA quick_check;").fetchone()
            except DATABASE_ERROR_TYPES as exc:
                raise DatabaseCorruptionError(self.path, str(exc)) from exc

        status = str(row[0]).strip() if row and row[0] is not None else "quick_check sem retorno"
        if status.lower() != "ok":
            raise DatabaseCorruptionError(self.path, status)

    def _wrap_database_error(self, exc):
        if isinstance(exc, DatabaseCorruptionError):
            return exc
        if _is_corruption_message(str(exc)):
            return DatabaseCorruptionError(self.path, str(exc))
        return exc

    def _configure_connection(self, conn: sqlite3.Connection) -> None:
        conn.execute("PRAGMA foreign_keys = ON;")
        conn.execute(f"PRAGMA busy_timeout = {SQLITE_BUSY_TIMEOUT_MS};")
        conn.execute("PRAGMA journal_mode = WAL;")
        conn.execute("PRAGMA synchronous = NORMAL;")
        conn.execute("PRAGMA temp_store = MEMORY;")
        conn.execute("PRAGMA wal_autocheckpoint = 1000;")

    def is_connection_usable(self) -> bool:
        """Validate the shared connection before it is reused by the web app."""
        with self._lock:
            conn = getattr(self, "conn", None)
            if conn is None:
                return False

            raw = getattr(conn, "raw", conn)
            if bool(getattr(conn, "_closed", False)) or bool(getattr(raw, "closed", False)):
                return False
            if bool(getattr(raw, "broken", False)):
                return False

            try:
                conn.execute("SELECT 1").fetchone()
                if DB_ENGINE in {"postgres", "postgresql"}:
                    # psycopg starts a transaction even for SELECT. Leaving this
                    # ping open can trigger idle_in_transaction_session_timeout.
                    conn.rollback()
                return True
            except DATABASE_ERROR_TYPES:
                return False

    def _fetchone(self, sql: str, params: Iterable = ()) -> sqlite3.Row | None:
        with self._lock:
            try:
                return self.conn.execute(sql, tuple(params)).fetchone()
            except DATABASE_ERROR_TYPES as exc:
                raise self._wrap_database_error(exc) from exc

    def _fetchall(self, sql: str, params: Iterable = ()) -> list[sqlite3.Row]:
        with self._lock:
            try:
                return self.conn.execute(sql, tuple(params)).fetchall()
            except DATABASE_ERROR_TYPES as exc:
                raise self._wrap_database_error(exc) from exc

    def _execute(self, sql: str, params: Iterable = (), commit: bool = False) -> sqlite3.Cursor:
        with self._lock:
            try:
                cursor = self.conn.execute(sql, tuple(params))
                if commit:
                    self.conn.commit()
                return cursor
            except DATABASE_ERROR_TYPES as exc:
                raise self._wrap_database_error(exc) from exc

    def _executemany(self, sql: str, params: Iterable[Iterable], commit: bool = False) -> sqlite3.Cursor:
        with self._lock:
            try:
                cursor = self.conn.executemany(sql, params)
                if commit:
                    self.conn.commit()
                return cursor
            except DATABASE_ERROR_TYPES as exc:
                raise self._wrap_database_error(exc) from exc

    def _table_config(self, tabela: str) -> dict[str, str]:
        if tabela not in REFERENCE_TABLES:
            raise ValueError(f"Tabela de referencia invalida: {tabela}")
        return REFERENCE_TABLES[tabela]

    def _criar_tabelas(self) -> None:
        with self._lock:
            cur = self.conn.cursor()
            cur.executescript(
                """
                CREATE TABLE IF NOT EXISTS motoristas (
                    codigo INTEGER PRIMARY KEY,
                    nome   TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS fazendas (
                    codigo TEXT PRIMARY KEY,
                    nome   TEXT NOT NULL


                );
                CREATE TABLE IF NOT EXISTS talhoes (
                    id                      TEXT PRIMARY KEY,
                    fazenda_id_mestre       TEXT,
                    fazenda_codigo          TEXT NOT NULL,
                    fazenda_codigo_mestre   TEXT,
                    fazenda_nome            TEXT,
                    codigo                  TEXT NOT NULL,
                    nome                    TEXT,
                    area_ha                 REAL,
                    area_alq                REAL,
                    area_plantada_ha        REAL,
                    safra                   TEXT,
                    tipo_area               TEXT,
                    ativo                   INTEGER DEFAULT 1,
                    fonte                   TEXT NOT NULL DEFAULT 'balanca',
                    sincronizado_em         TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    UNIQUE(fazenda_codigo, codigo, fonte),
                    FOREIGN KEY(fazenda_codigo) REFERENCES fazendas(codigo)
                );
                CREATE TABLE IF NOT EXISTS sincronizacoes_cadastros (
                    fonte                TEXT PRIMARY KEY,
                    sincronizado_em      TEXT NOT NULL,
                    fazendas_lidas       INTEGER NOT NULL DEFAULT 0,
                    fazendas_alteradas   INTEGER NOT NULL DEFAULT 0,
                    talhoes_lidos        INTEGER NOT NULL DEFAULT 0,
                    talhoes_alterados    INTEGER NOT NULL DEFAULT 0,
                    origem               TEXT,
                    detalhes             TEXT
                );
                CREATE TABLE IF NOT EXISTS variedades (
                    id     INTEGER PRIMARY KEY AUTOINCREMENT,
                    nome   TEXT NOT NULL UNIQUE
                );
                CREATE TABLE IF NOT EXISTS notas (
                    numero           INTEGER PRIMARY KEY,
                    motorista_cod    INTEGER,
                    motorista_nome   TEXT,
                    caminhao         TEXT,
                    operador_cod     INTEGER,
                    operador_nome    TEXT,
                    colhedora        TEXT,
                    faz_muda_cod     TEXT,
                    faz_muda_nome    TEXT,
                    talhao           TEXT,
                    faz_plantio_cod  TEXT,
                    faz_plantio_nome TEXT,
                    variedade_id     INTEGER,
                    variedade_nome   TEXT,
                    data_colheita    TEXT,
                    data_plantio     TEXT,
                    duplicado        INTEGER DEFAULT 0
                );

                CREATE TABLE IF NOT EXISTS auditoria_correcoes (
                    id                  INTEGER PRIMARY KEY AUTOINCREMENT,
                    aplicado_em         TEXT NOT NULL,
                    nota_numero         INTEGER NOT NULL,
                    acao                TEXT NOT NULL,
                    motivo              TEXT,
                    data_colheita_ant   TEXT,
                    data_colheita_nova  TEXT,
                    data_plantio_ant    TEXT,
                    data_plantio_nova   TEXT,
                    backup_path         TEXT
                );

                CREATE TABLE IF NOT EXISTS notas_referencias_auditoria (
                    id                  INTEGER PRIMARY KEY AUTOINCREMENT,
                    nota_numero         INTEGER NOT NULL,
                    registrado_em       TEXT NOT NULL,
                    evento              TEXT NOT NULL,
                    referencias_json    TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS cidades (
                    id   INTEGER PRIMARY KEY AUTOINCREMENT,
                    nome TEXT NOT NULL


                );

                CREATE TABLE IF NOT EXISTS frota (
                    id     INTEGER PRIMARY KEY AUTOINCREMENT,
                    numero TEXT NOT NULL UNIQUE,
                    placa  TEXT,
                    status TEXT DEFAULT 'ATIVO'
                );

                CREATE TABLE IF NOT EXISTS frentes (
                    id   INTEGER PRIMARY KEY AUTOINCREMENT,
                    nome TEXT NOT NULL UNIQUE
                );

                CREATE TABLE IF NOT EXISTS localizacao_frentes (
                    id          INTEGER PRIMARY KEY AUTOINCREMENT,
                    data        TEXT NOT NULL,
                    frente_id   INTEGER NOT NULL,
                    fazenda_cod TEXT NOT NULL,
                    FOREIGN KEY(frente_id) REFERENCES frentes(id),
                    FOREIGN KEY(fazenda_cod) REFERENCES fazendas(codigo)
                );

                CREATE TABLE IF NOT EXISTS escala_viagem (
                    id         INTEGER PRIMARY KEY AUTOINCREMENT,
                    onibus_id  INTEGER NOT NULL,
                    frente_id  INTEGER NOT NULL,
                    cidade_id  INTEGER NOT NULL,
                    hora_ida   TEXT NOT NULL,
                    hora_volta TEXT NOT NULL,
                    FOREIGN KEY(onibus_id) REFERENCES frota(id),
                    FOREIGN KEY(frente_id) REFERENCES frentes(id),
                    FOREIGN KEY(cidade_id) REFERENCES cidades(id)
                );
                """
            )
            self.conn.commit()

    def _table_columns(self, table_name: str) -> set[str]:
        with self._lock:
            rows = self.conn.execute(f"PRAGMA table_info({table_name})").fetchall()
        return {str(row["name"]) for row in rows}

    def _ensure_column(self, table_name: str, column_name: str, definition: str) -> None:
        if column_name in self._table_columns(table_name):
            return
        with self._lock:
            self.conn.execute(f"ALTER TABLE {table_name} ADD COLUMN {column_name} {definition}")
            self.conn.commit()

    def _migrar_schema(self) -> None:
        self._ensure_column("fazendas", "id_mestre", "TEXT")
        self._ensure_column("fazendas", "codigo_mestre", "TEXT")
        self._ensure_column("fazendas", "fonte_mestre", "TEXT")
        self._ensure_column("fazendas", "sincronizado_em", "TEXT")

    def _criar_indices(self) -> None:
        with self._lock:
            self.conn.executescript(
                """
                CREATE INDEX IF NOT EXISTS idx_notas_data_colheita ON notas(data_colheita);
                CREATE INDEX IF NOT EXISTS idx_notas_faz_muda_cod ON notas(faz_muda_cod);
                CREATE INDEX IF NOT EXISTS idx_notas_faz_plantio_cod ON notas(faz_plantio_cod);
                CREATE INDEX IF NOT EXISTS idx_notas_data_colheita_faz_muda_cod ON notas(data_colheita, faz_muda_cod);
                CREATE INDEX IF NOT EXISTS idx_notas_data_colheita_faz_plantio_cod ON notas(data_colheita, faz_plantio_cod);
                CREATE INDEX IF NOT EXISTS idx_notas_motorista_nome ON notas(motorista_nome);
                CREATE INDEX IF NOT EXISTS idx_notas_operador_nome ON notas(operador_nome);
                CREATE INDEX IF NOT EXISTS idx_notas_faz_plantio_nome ON notas(faz_plantio_nome);
                CREATE INDEX IF NOT EXISTS idx_notas_faz_muda_nome ON notas(faz_muda_nome);
                CREATE INDEX IF NOT EXISTS idx_notas_variedade_nome ON notas(variedade_nome);
                CREATE INDEX IF NOT EXISTS idx_notas_talhao ON notas(talhao);
                CREATE INDEX IF NOT EXISTS idx_notas_caminhao ON notas(caminhao);
                CREATE INDEX IF NOT EXISTS idx_auditoria_correcoes_nota ON auditoria_correcoes(nota_numero, aplicado_em);
                CREATE INDEX IF NOT EXISTS idx_notas_referencias_auditoria_nota ON notas_referencias_auditoria(nota_numero, registrado_em);
                CREATE INDEX IF NOT EXISTS idx_fazendas_nome ON fazendas(nome);
                CREATE INDEX IF NOT EXISTS idx_fazendas_codigo_mestre ON fazendas(codigo_mestre);
                CREATE INDEX IF NOT EXISTS idx_talhoes_fazenda_codigo ON talhoes(fazenda_codigo);
                CREATE INDEX IF NOT EXISTS idx_talhoes_fazenda_codigo_talhao ON talhoes(fazenda_codigo, codigo);
                CREATE INDEX IF NOT EXISTS idx_talhoes_fonte ON talhoes(fonte);
                CREATE INDEX IF NOT EXISTS idx_motoristas_nome ON motoristas(nome);
                CREATE INDEX IF NOT EXISTS idx_variedades_nome ON variedades(nome);
                CREATE INDEX IF NOT EXISTS idx_localizacao_frentes_data ON localizacao_frentes(data);
                CREATE INDEX IF NOT EXISTS idx_localizacao_frentes_frente_data ON localizacao_frentes(frente_id, data);
                CREATE INDEX IF NOT EXISTS idx_localizacao_frentes_fazenda ON localizacao_frentes(fazenda_cod);
                """
            )
            self.conn.commit()

    def _seed_from_excel(self) -> None:
        if not EXCEL_ORIGEM.exists():
            LOGGER.info("Excel de origem nao encontrado em %s; seed inicial ignorado", EXCEL_ORIGEM)
            return
        if load_workbook is None:
            LOGGER.warning("openpyxl indisponivel; seed inicial do Excel foi ignorado")
            return

        wb = load_workbook(EXCEL_ORIGEM, data_only=True, keep_vba=True)
        try:
            if ABAS_EXCEL["motoristas"] in wb.sheetnames:
                ws = wb[ABAS_EXCEL["motoristas"]]
                lista = []
                for row in ws.iter_rows(min_row=2, max_col=2):
                    try:
                        if row[0].value and row[1].value:
                            lista.append((int(row[0].value), str(row[1].value).strip()))
                    except (TypeError, ValueError):
                        continue
                if lista:
                    self._executemany(
                        "INSERT OR IGNORE INTO motoristas(codigo, nome) VALUES (?, ?)",
                        lista,
                        commit=True
                    )

            if ABAS_EXCEL["fazendas"] in wb.sheetnames:
                ws = wb[ABAS_EXCEL["fazendas"]]
                lista = []
                for row in ws.iter_rows(min_row=2, max_col=2):
                    if row[0].value and row[1].value:
                        lista.append((str(row[0].value).strip(), str(row[1].value).strip()))
                if lista:
                    self._executemany(
                        "INSERT OR IGNORE INTO fazendas(codigo, nome) VALUES (?, ?)",
                        lista,
                        commit=True
                    )

            if ABAS_EXCEL["variedades"] in wb.sheetnames:
                ws = wb[ABAS_EXCEL["variedades"]]
                lista = [(str(cell.value).strip(),) for cell in ws["A"][1:] if cell.value]
                if lista:
                    self._executemany(
                        "INSERT OR IGNORE INTO variedades(nome) VALUES (?)",
                        lista,
                        commit=True
                    )
        finally:
            wb.close()

    def _listar_referencias_local(
        self,
        tabela: str,
        q: str = "",
        limit: int | None = None
    ) -> list[sqlite3.Row]:
        config = self._table_config(tabela)
        id_col = config["id_col"]
        sql = f"SELECT {id_col}, nome FROM {tabela}"
        params: list[object] = []
        q_txt = str(q or "").strip()

        if q_txt:
            sql += f" WHERE CAST({id_col} AS TEXT) LIKE ? OR UPPER(nome) LIKE UPPER(?)"
            params.extend((f"%{q_txt}%", f"%{q_txt}%"))

        sql += " ORDER BY nome"
        if limit is not None:
            sql += " LIMIT ?"
            params.append(max(0, int(limit)))

        return self._fetchall(sql, tuple(params))

    def _colaboradores_provider(self):
        """Return (provider, authoritative) without silently downgrading errors."""
        if get_colaboradores_provider is None:
            return None, False
        try:
            provider = get_colaboradores_provider()
            if provider is None:
                return None, False
            return provider, bool(getattr(provider, "enabled", True))
        except Exception as exc:
            LOGGER.warning("Falha ao inicializar referencias do Portal Colaboradores: %s", exc)
            return None, True

    def _listar_motoristas_colaboradores(self, q: str = "", limit: int | None = None) -> tuple[bool, list]:
        provider, authoritative = self._colaboradores_provider()
        if not authoritative or provider is None:
            return authoritative, []
        try:
            return True, provider.listar_referencias(q=q, limit=limit)
        except Exception as exc:
            LOGGER.warning("Falha ao consultar motoristas no Portal Colaboradores: %s", exc)
            return True, []

    def _buscar_motorista_colaboradores(self, valor) -> tuple[bool, dict | None]:
        provider, authoritative = self._colaboradores_provider()
        if not authoritative or provider is None:
            return authoritative, None
        try:
            return True, provider.buscar_referencia(valor)
        except Exception as exc:
            LOGGER.warning("Falha ao resolver motorista no Portal Colaboradores: %s", exc)
            return True, None

    @staticmethod
    def _dedupe_referencias_por_codigo(rows: list, limit: int | None = None) -> list:
        deduped = []
        vistos = set()
        for row in rows:
            codigo = str(row["codigo"] if row["codigo"] is not None else "").strip()
            chave = codigo.upper()
            if chave in vistos:
                continue
            vistos.add(chave)
            deduped.append(row)
        if limit is not None:
            return deduped[: max(0, int(limit))]
        return deduped

    def listar_referencias_filtradas(
        self,
        tabela: str,
        q: str = "",
        limit: int | None = None
    ) -> list:
        if tabela == "motoristas":
            authoritative, colaboradores = self._listar_motoristas_colaboradores(q=q, limit=limit)
            if authoritative:
                return self._dedupe_referencias_por_codigo(colaboradores, limit=limit)
            locais = self._listar_referencias_local(tabela, q=q, limit=limit)
            return self._dedupe_referencias_por_codigo(locais, limit=limit)
        return self._listar_referencias_local(tabela, q=q, limit=limit)

    def listar_referencias(self, tabela: str) -> list:
        return self.listar_referencias_filtradas(tabela)

    def listar_motoristas(self) -> list[sqlite3.Row]:
        return self.listar_referencias("motoristas")

    def listar_fazendas(self) -> list[sqlite3.Row]:
        return self.listar_referencias("fazendas")

    def listar_talhoes(self, fazenda_codigo: str | None = None) -> list[sqlite3.Row]:
        return self.listar_talhoes_filtrados(fazenda_codigo=fazenda_codigo)

    def listar_talhoes_filtrados(
        self,
        fazenda_codigo: str | None = None,
        q: str = "",
        limit: int | None = None
    ) -> list[sqlite3.Row]:
        filtros = ["ativo = 1"]
        params: list[object] = []
        if fazenda_codigo:
            codigo = str(fazenda_codigo).strip()
            filtros.append("(fazenda_codigo = ? OR fazenda_codigo_mestre = ?)")
            params.extend((codigo, codigo))
        q_txt = str(q or "").strip()
        if q_txt:
            filtros.append(
                """
                (
                    CAST(codigo AS TEXT) LIKE ?
                    OR UPPER(COALESCE(nome, '')) LIKE UPPER(?)
                    OR UPPER(COALESCE(fazenda_nome, '')) LIKE UPPER(?)
                )
                """
            )
            like = f"%{q_txt}%"
            params.extend((like, like, like))

        sql = """
            SELECT *
            FROM talhoes
            WHERE {where_clause}
            ORDER BY fazenda_nome, {codigo_order}, codigo
        """.format(where_clause=" AND ".join(filtros), codigo_order=CODIGO_NUMERIC_ORDER)
        if limit is not None:
            sql += " LIMIT ?"
            params.append(max(0, int(limit)))
        return self._fetchall(sql, tuple(params))

    def buscar_talhao(self, valor, fazenda_codigo: str | None = None) -> sqlite3.Row | None:
        valor_txt = str(valor or "").strip()
        if not valor_txt:
            return None
        if " - " in valor_txt:
            codigo, nome = valor_txt.split(" - ", 1)
            return self.buscar_talhao(codigo, fazenda_codigo=fazenda_codigo) or self.buscar_talhao(
                nome,
                fazenda_codigo=fazenda_codigo
            )

        filtros = ["ativo = 1", "(codigo = ? OR UPPER(TRIM(COALESCE(nome, ''))) = UPPER(TRIM(?)))"]
        params: list[object] = [valor_txt, valor_txt]
        if fazenda_codigo:
            codigo = str(fazenda_codigo).strip()
            filtros.append("(fazenda_codigo = ? OR fazenda_codigo_mestre = ?)")
            params.extend((codigo, codigo))
        return self._fetchone(
            """
            SELECT *
            FROM talhoes
            WHERE {where_clause}
            ORDER BY fazenda_nome, {codigo_order}, codigo
            LIMIT 1
            """.format(where_clause=" AND ".join(filtros), codigo_order=CODIGO_NUMERIC_ORDER),
            tuple(params)
        )

    def listar_variedades(self) -> list[sqlite3.Row]:
        return self.listar_referencias("variedades")


    @staticmethod
    def _formatar_codigo_nome(codigo, nome) -> str:
        codigo_txt = "" if codigo in (None, "") else str(codigo)
        nome_txt = "" if nome in (None, "") else str(nome)
        if codigo_txt and nome_txt:
            return f"{codigo_txt} - {nome_txt}"
        return codigo_txt or nome_txt

    def buscar_por_codigo(self, tabela: str, codigo) -> dict | None:
        config = self._table_config(tabela)
        if codigo in (None, ""):
            return None

        query = f"SELECT {config['id_col']} AS codigo, nome FROM {tabela} WHERE {config['id_col']} = ?"
        row = self._fetchone(query, (codigo,))
        if row:
            return dict(row)

        if tabela == "fazendas" and "-" not in str(codigo) and len(str(codigo)) >= 4:
            codigo_txt = str(codigo)
            codigo_formatado = codigo_txt[:3] + "-" + codigo_txt[3:]
            row = self._fetchone(query, (codigo_formatado,))
            if row:
                return dict(row)
        return None

    def buscar_referencia(self, tabela: str, valor, col_id: str | None = None) -> sqlite3.Row | None:
        if valor in (None, ""):
            return None

        config = self._table_config(tabela)
        campo_id = col_id or config["id_col"]
        valor_txt = str(valor).strip()

        if tabela == "motoristas":
            authoritative, row_colaboradores = self._buscar_motorista_colaboradores(valor_txt)
            if row_colaboradores:
                return row_colaboradores
            if authoritative:
                return None

        row = self._fetchone(
            f"SELECT * FROM {tabela} WHERE {campo_id} = ?",
            (valor_txt,)
        )
        if row:
            return row

        if tabela == "fazendas" and campo_id == "codigo" and "-" not in valor_txt and len(valor_txt) >= 4:
            codigo_formatado = valor_txt[:3] + "-" + valor_txt[3:]
            row = self._fetchone(
                f"SELECT * FROM {tabela} WHERE {campo_id} = ?",
                (codigo_formatado,)
            )
            if row:
                return row

        return self._fetchone(
            f"SELECT * FROM {tabela} WHERE UPPER(TRIM(nome)) = UPPER(TRIM(?))",
            (valor_txt,)
        )

    def resolver_referencia(self, tabela: str, texto: str, col_id: str | None = None) -> sqlite3.Row | None:
        valor = (texto or "").strip()
        if not valor:
            return None

        if " - " in valor:
            codigo, nome = valor.split(" - ", 1)
            return self.buscar_referencia(tabela, codigo.strip(), col_id=col_id) or self.buscar_referencia(
                tabela,
                nome.strip(),
                col_id=col_id
            )

        return self.buscar_referencia(tabela, valor, col_id=col_id)

    def buscar_nota(self, numero) -> sqlite3.Row | None:
        return self._fetchone("SELECT * FROM notas WHERE numero = ?", (numero,))

    def gerar_numero_duplicado(self, numero_base: int) -> int:
        base = int(numero_base)
        if base <= 0:
            raise ValueError("Numero base invalido para duplicacao.")

        for sequencia in range(1, 10000):
            candidato = int(f"{base}{sequencia:02d}")
            if not self.buscar_nota(candidato):
                return candidato

        raise RuntimeError(f"Nao foi possivel gerar um numero duplicado livre para a nota {base}.")

    def listar_notas_por_numeros(
        self, numeros: Iterable[int], *, lock_for_update: bool = False
    ) -> list[sqlite3.Row]:
        numeros_validos = sorted({int(numero) for numero in numeros})
        if not numeros_validos:
            return []

        placeholders = ", ".join("?" for _ in numeros_validos)
        lock_clause = (
            " FOR UPDATE"
            if lock_for_update and DB_ENGINE in {"postgres", "postgresql"}
            else ""
        )
        return self._fetchall(
            f"""
            SELECT numero, faz_muda_cod, faz_muda_nome, faz_plantio_cod, faz_plantio_nome,
                   data_colheita, data_plantio
            FROM notas
            WHERE numero IN ({placeholders})
            ORDER BY numero{lock_clause}
            """,
            tuple(numeros_validos)
        )

    def _resolver_datas_correcao(self, row: sqlite3.Row, acao: str, nova_data: str | None) -> tuple[str | None, str | None]:
        data_colheita_atual = row["data_colheita"]
        data_plantio_atual = row["data_plantio"]

        if acao == "colheita_para_plantio":
            return data_plantio_atual, data_plantio_atual

        if acao == "definir_colheita":
            data_sql = _normalize_note_date(nova_data, "Nova data de colheita")
            return data_sql, data_plantio_atual

        if acao == "definir_plantio":
            data_sql = _normalize_note_date(nova_data, "Nova data de plantio")
            return data_colheita_atual, data_sql

        if acao == "definir_ambas":
            data_sql = _normalize_note_date(nova_data, "Nova data")
            return data_sql, data_sql

        raise ValueError(f"Acao de correcao invalida: {acao}")

    def preview_correcao_notas(
        self,
        numeros: Iterable[int],
        acao: str,
        nova_data: str | None = None,
        *,
        lock_for_update: bool = False
    ) -> dict[str, object]:
        numeros_validos = sorted({int(numero) for numero in numeros})
        rows = self.listar_notas_por_numeros(
            numeros_validos, lock_for_update=lock_for_update
        )
        rows_por_numero = {int(row["numero"]): row for row in rows}
        preview = []
        faltantes = []

        for numero in numeros_validos:
            row = rows_por_numero.get(numero)
            if row is None:
                faltantes.append(numero)
                continue

            nova_colheita, novo_plantio = self._resolver_datas_correcao(row, acao, nova_data)
            preview.append(
                {
                    "numero": int(row["numero"]),
                    "faz_muda": self._formatar_codigo_nome(row["faz_muda_cod"], row["faz_muda_nome"]),
                    "faz_plantio": self._formatar_codigo_nome(row["faz_plantio_cod"], row["faz_plantio_nome"]),
                    "data_colheita_atual": row["data_colheita"],
                    "data_colheita_nova": nova_colheita,
                    "data_plantio_atual": row["data_plantio"],
                    "data_plantio_nova": novo_plantio,
                    "alterado": (
                        nova_colheita != row["data_colheita"]
                        or novo_plantio != row["data_plantio"]
                    ),
                }
            )

        token_payload = {
            "acao": acao,
            "nova_data": nova_data,
            "preview": preview,
            "faltantes": faltantes,
        }
        preview_token = hashlib.sha256(
            json.dumps(token_payload, ensure_ascii=False, sort_keys=True, default=str).encode("utf-8")
        ).hexdigest()
        return {"preview": preview, "faltantes": faltantes, "preview_token": preview_token}

    def aplicar_correcao_notas(
        self,
        numeros: Iterable[int],
        acao: str,
        nova_data: str | None = None,
        motivo: str = "",
        backup_path: str | Path | None = None,
        expected_preview_token: str | None = None
    ) -> dict[str, object]:
        with self._lock:
            try:
                # End any read-only transaction left by the preview endpoint,
                # then lock the correction snapshot through commit.
                self.conn.rollback()
                if DB_ENGINE not in {"postgres", "postgresql"}:
                    self.conn.execute("BEGIN IMMEDIATE")
                resultado_preview = self.preview_correcao_notas(
                    numeros, acao, nova_data, lock_for_update=True
                )
                preview = resultado_preview["preview"]
                faltantes = resultado_preview["faltantes"]
                alteracoes = [item for item in preview if item["alterado"]]

                if expected_preview_token and expected_preview_token != resultado_preview["preview_token"]:
                    raise CorrectionPreviewConflict(
                        "As notas mudaram depois da prévia. Gere uma nova prévia antes de aplicar."
                    )
                if not preview:
                    raise ValueError("Nenhuma nota encontrada para a correcao solicitada.")
                if not alteracoes:
                    raise ValueError("Nenhuma data seria alterada com a correcao informada.")

                timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
                destino_backup = Path(backup_path) if backup_path else APP_ROOT / "backups" / f"backup_correcao_{timestamp}.db"
                destino_backup = self.create_backup(destino_backup)

                aplicado_em = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                motivo_txt = (motivo or "").strip()
                for item in alteracoes:
                    self.conn.execute(
                        """
                        UPDATE notas
                        SET data_colheita = ?, data_plantio = ?
                        WHERE numero = ?
                        """,
                        (
                            item["data_colheita_nova"],
                            item["data_plantio_nova"],
                            item["numero"]
                        )
                    )
                    self.conn.execute(
                        """
                        INSERT INTO auditoria_correcoes (
                            aplicado_em, nota_numero, acao, motivo,
                            data_colheita_ant, data_colheita_nova,
                            data_plantio_ant, data_plantio_nova,
                            backup_path
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            aplicado_em,
                            item["numero"],
                            acao,
                            motivo_txt,
                            item["data_colheita_atual"],
                            item["data_colheita_nova"],
                            item["data_plantio_atual"],
                            item["data_plantio_nova"],
                            str(destino_backup)
                        )
                    )
                self.conn.commit()
            except Exception as exc:
                self.conn.rollback()
                if isinstance(exc, DATABASE_ERROR_TYPES):
                    raise self._wrap_database_error(exc) from exc
                raise

        return {
            "quantidade_alterada": len(alteracoes),
            "quantidade_faltante": len(faltantes),
            "faltantes": faltantes,
            "backup_path": str(destino_backup),
        }

    def listar_logs_correcao(self, limit: int = 50) -> list[sqlite3.Row]:
        return self._fetchall(
            """
            SELECT aplicado_em, nota_numero, acao, motivo,
                   data_colheita_ant, data_colheita_nova,
                   data_plantio_ant, data_plantio_nova,
                   backup_path
            FROM auditoria_correcoes
            ORDER BY id DESC
            LIMIT ?
            """,
            (int(limit),)
        )

    def montar_nota_payload(self, dados: dict) -> dict:
        payload = {campo: dados.get(campo) for campo in NOTA_FIELDS}
        payload["numero"] = int(payload["numero"])
        payload["data_colheita"] = _normalize_note_date(
            payload.get("data_colheita"),
            "Data de colheita"
        )
        payload["data_plantio"] = _normalize_note_date(
            payload.get("data_plantio"),
            "Data de plantio"
        )
        payload["duplicado"] = int(bool(dados.get("duplicado", 0)))
        return payload

    def inserir_nota(self, dados: dict, force: bool = False) -> None:
        payload = self.montar_nota_payload(dados)
        params = [payload[campo] for campo in NOTA_FIELDS] + [payload["duplicado"]]
        columns = ", ".join((*NOTA_FIELDS, "duplicado"))
        placeholders = ", ".join("?" for _ in range(len(NOTA_FIELDS) + 1))
        conflict_sql = ""
        if force:
            updated_columns = (*NOTA_FIELDS[1:], "duplicado")
            assignments = ", ".join(f"{column}=excluded.{column}" for column in updated_columns)
            conflict_sql = f" ON CONFLICT(numero) DO UPDATE SET {assignments}"

        with self._lock:
            try:
                self.conn.execute(
                    f"INSERT INTO notas ({columns}) VALUES ({placeholders}){conflict_sql}",
                    params
                )
                self.registrar_referencias_nota(
                    payload["numero"],
                    dados.get("_referencias_origem"),
                    evento=str(dados.get("_audit_event") or "nota.save"),
                    commit=False
                )
                self.conn.commit()
            except DATABASE_ERROR_TYPES as exc:
                self.conn.rollback()
                raise self._wrap_database_error(exc) from exc

    def registrar_referencias_nota(
        self,
        nota_numero: int,
        referencias: dict | None,
        evento: str = "nota.save",
        commit: bool = True
    ) -> None:
        if not referencias:
            return
        registrado_em = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        payload = json.dumps(referencias, ensure_ascii=False, sort_keys=True)
        with self._lock:
            self.conn.execute(
                """
                INSERT INTO notas_referencias_auditoria (
                    nota_numero, registrado_em, evento, referencias_json
                ) VALUES (?, ?, ?, ?)
                """,
                (int(nota_numero), registrado_em, evento, payload)
            )
            if commit:
                self.conn.commit()

    def listar_referencias_auditoria(self, nota_numero: int, limit: int = 20) -> list[dict]:
        rows = self._fetchall(
            """
            SELECT nota_numero, registrado_em, evento, referencias_json
            FROM notas_referencias_auditoria
            WHERE nota_numero = ?
            ORDER BY id DESC
            LIMIT ?
            """,
            (int(nota_numero), int(limit))
        )
        return [
            {
                "nota_numero": row["nota_numero"],
                "registrado_em": row["registrado_em"],
                "evento": row["evento"],
                "referencias": json.loads(row["referencias_json"] or "{}"),
            }
            for row in rows
        ]

    def listar_notas(self, data_inicio: str | None = None, data_fim: str | None = None) -> list[sqlite3.Row]:
        return self.buscar_notas_historico(data_inicio=data_inicio, data_fim=data_fim)

    def buscar_notas_historico(
        self,
        data_inicio: str | None = None,
        data_fim: str | None = None,
        numero_prefixo: str | None = None
    ) -> list[sqlite3.Row]:
        sql = """
            SELECT numero, motorista_cod, motorista_nome, caminhao, operador_cod,
                   operador_nome, colhedora, faz_muda_cod, faz_muda_nome, talhao,
                   faz_plantio_cod, faz_plantio_nome, variedade_nome, data_colheita,
                   data_plantio
            FROM notas
        """
        filtros: list[str] = []
        params: list[str] = []

        if data_inicio and data_fim:
            filtros.append("data_colheita BETWEEN ? AND ?")
            params.extend((data_inicio, data_fim))
        if numero_prefixo:
            prefixo = str(numero_prefixo).strip()
            if not prefixo.isdigit():
                return []

            faixas = []
            base = int(prefixo)
            max_digitos = 9
            digitos_prefixo = len(prefixo)
            for extra_digitos in range(max(0, max_digitos - digitos_prefixo) + 1):
                fator = 10**extra_digitos
                inicio = base * fator
                fim = inicio + fator - 1
                faixas.append("(numero BETWEEN ? AND ?)")
                params.extend((str(inicio), str(fim)))
            filtros.append("(" + " OR ".join(faixas) + ")")
        if filtros:
            sql += " WHERE " + " AND ".join(filtros)
        sql += " ORDER BY numero DESC"
        return self._fetchall(sql, params)

    def excluir_nota(self, numero) -> None:
        self._execute("DELETE FROM notas WHERE numero = ?", (numero,), commit=True)

    def contar_notas_total(self) -> int:
        row = self._fetchone("SELECT COUNT(*) AS total FROM notas")
        return int(row["total"] if row else 0)

    def contar_notas_data(self, data_sql: str) -> int:
        row = self._fetchone(
            "SELECT COUNT(*) AS total FROM notas WHERE data_colheita = ?",
            (data_sql,)
        )
        return int(row["total"] if row else 0)

    def contar_notas_periodo(self, d_ini: str, d_fim: str) -> int:
        row = self._fetchone(
            "SELECT COUNT(*) AS total FROM notas WHERE data_colheita BETWEEN ? AND ?",
            (d_ini, d_fim)
        )
        return int(row["total"] if row else 0)

    def contar_dias_ativos_periodo(self, d_ini: str, d_fim: str) -> int:
        row = self._fetchone(
            """
            SELECT COUNT(DISTINCT data_colheita) AS total
            FROM notas
            WHERE data_colheita BETWEEN ? AND ?
            """,
            (d_ini, d_fim)
        )
        return int(row["total"] if row else 0)

    def top_por_coluna(self, coluna: str, d_ini: str, d_fim: str, limit: int = 5) -> list:
        allowed_cols = {
            "motorista_nome",
            "operador_nome",
            "colhedora",
            "variedade_nome",
            "faz_muda_nome",
            "faz_plantio_nome",
        }
        if coluna not in allowed_cols:
            raise ValueError(f"Coluna nao permitida para ranking: {coluna}")

        fazenda_cod_cols = {
            "faz_muda_nome": "faz_muda_cod",
            "faz_plantio_nome": "faz_plantio_cod",
        }
        if coluna in fazenda_cod_cols:
            codigo_col = fazenda_cod_cols[coluna]
            rows = self._fetchall(
                f"""
                SELECT {codigo_col} AS codigo, {coluna} AS nome
                FROM notas
                WHERE data_colheita BETWEEN ? AND ?
                  AND (
                    ({coluna} IS NOT NULL AND TRIM({coluna}) <> '')
                    OR ({codigo_col} IS NOT NULL AND TRIM({codigo_col}) <> '')
                  )
                """,
                (d_ini, d_fim)
            )

            grupos: dict[str, dict] = {}
            for row in rows:
                codigo = str(row["codigo"] or "").strip()
                nome = str(row["nome"] or "").strip()
                if not codigo and not nome:
                    continue
                chave = codigo or f"NOME:{nome.upper()}"
                if chave not in grupos:
                    grupos[chave] = {"codigo": codigo, "nomes": Counter(), "qtd": 0}

                grupos[chave]["nomes"][nome or codigo or "-"] += 1
                grupos[chave]["qtd"] += 1

            ranking = []
            for grupo in grupos.values():
                nome = grupo["nomes"].most_common(1)[0][0]
                codigo = grupo["codigo"]
                ranking.append(
                    {
                        "codigo": codigo,
                        "nome_fazenda": nome,
                        "nome": self._formatar_codigo_nome(codigo, nome) or "-",
                        "qtd": grupo["qtd"],
                    }
                )

            ranking.sort(key=lambda item: (-item["qtd"], str(item["nome"]).upper()))
            return ranking[: int(limit)]

        return self._fetchall(
            f"""
            SELECT {coluna} AS nome, COUNT(*) AS qtd
            FROM notas
            WHERE data_colheita BETWEEN ? AND ?
              AND {coluna} IS NOT NULL
              AND TRIM({coluna}) <> ''
            GROUP BY {coluna}
            ORDER BY qtd DESC, {coluna}
            LIMIT ?
            """,
            (d_ini, d_fim, int(limit))
        )

    def top_motorista_do_dia(self, data_sql: str) -> sqlite3.Row | None:
        return self._fetchone(
            """
            SELECT motorista_nome, COUNT(*) AS qtd
            FROM notas
            WHERE data_colheita = ?
              AND motorista_nome IS NOT NULL
              AND TRIM(motorista_nome) <> ''
            GROUP BY motorista_nome
            ORDER BY qtd DESC, motorista_nome
            LIMIT 1
            """,
            (data_sql,)
        )

    def dataframe_por_query(self, sql: str, params: Iterable = ()):
        import pandas as pd

        if DB_ENGINE in {"postgres", "postgresql"}:
            return pd.read_sql_query(translate_sql(sql), self.conn.raw, params=tuple(params))

        conn = sqlite3.connect(self.path)
        try:
            return pd.read_sql_query(sql, conn, params=tuple(params))
        finally:
            conn.close()

    def dataframe_historico(self, data_inicio: str | None = None, data_fim: str | None = None):
        sql = "SELECT * FROM notas"
        params: tuple[str, str] | tuple[()] = ()
        if data_inicio and data_fim:
            sql += " WHERE data_colheita BETWEEN ? AND ?"
            params = (data_inicio, data_fim)
        sql += " ORDER BY numero DESC"
        return self.dataframe_por_query(sql, params)

    def dataframe_fluxo(self, d_ini: str, d_fim: str):
        return self.dataframe_por_query(
            """
            SELECT data_colheita AS Data, faz_muda_nome AS Origem,
                   talhao AS Talhao, variedade_nome AS Variedade,
                   faz_plantio_nome AS Destino, COUNT(*) AS Qtd
            FROM notas
            WHERE data_colheita BETWEEN ? AND ?
            GROUP BY data_colheita, faz_muda_nome, talhao, variedade_nome, faz_plantio_nome
            ORDER BY data_colheita
            """,
            (d_ini, d_fim)
        )

    def cadastrar_novo(self, tipo: str, codigo, nome: str):
        try:
            if tipo == "motorista":
                self.adicionar_motorista(codigo, nome)
            elif tipo == "fazenda":
                self.adicionar_fazenda(codigo, nome)
            elif tipo == "variedade":
                self.adicionar_variedade(nome)
            else:
                raise ValueError(f"Tipo de cadastro invalido: {tipo}")
            return True, "Sucesso"
        except Exception as exc:
            return False, str(exc)

    def adicionar_motorista(self, codigo, nome: str) -> None:
        self._execute(
            "INSERT INTO motoristas (codigo, nome) VALUES (?, ?)",
            (int(codigo), str(nome).strip()),
            commit=True
        )

    def adicionar_fazenda(self, codigo, nome: str) -> None:
        self._execute(
            "INSERT INTO fazendas (codigo, nome) VALUES (?, ?)",
            (str(codigo).strip(), str(nome).strip()),
            commit=True
        )

    def adicionar_variedade(self, nome: str) -> None:
        self._execute(
            "INSERT INTO variedades (nome) VALUES (?)",
            (str(nome).strip(),),
            commit=True
        )

    def excluir_cadastro(self, tabela: str, valor_id) -> None:
        config = self._table_config(tabela)
        self._execute(
            f"DELETE FROM {tabela} WHERE {config['id_col']} = ?",
            (valor_id,),
            commit=True
        )

    def create_backup(self, destination_path: str | Path) -> Path:
        destination = Path(destination_path)
        destination.parent.mkdir(parents=True, exist_ok=True)

        if DB_ENGINE in {"postgres", "postgresql"}:
            if destination.suffix.lower() == ".db":
                destination = destination.with_suffix(".dump")
            run_pg_dump(DATABASE_URL, destination)
            return destination

        with self._lock:
            source = sqlite3.connect(self.path)
            dst = sqlite3.connect(destination)
            try:
                source.backup(dst)
                dst.commit()
            finally:
                dst.close()
                source.close()

        return destination

    def restore_from_backup(self, backup_path: str | Path) -> None:
        origem = Path(backup_path)
        if not origem.exists():
            raise FileNotFoundError(f"Backup nao encontrado: {origem}")

        if DB_ENGINE in {"postgres", "postgresql"}:
            raise RuntimeError("Restore PostgreSQL do Sistema de Notas deve ser feito pelo procedimento operacional de banco.")

        try:
            with self._lock:
                self.conn.close()
                restore_sqlite_backup(self.path, origem)
                self.conn = self._open_connection()
                self._validate_integrity()
        except DATABASE_ERROR_TYPES as exc:
            raise self._wrap_database_error(exc) from exc


    def resumo_cadastros_mestre(self) -> dict[str, int | str | None]:
        fazendas = self._fetchone(
            "SELECT COUNT(*) AS total FROM fazendas WHERE fonte_mestre = 'balanca'"
        )
        talhoes = self._fetchone("SELECT COUNT(*) AS total FROM talhoes WHERE fonte = 'balanca'")
        sync = self._fetchone(
            """
            SELECT sincronizado_em
            FROM sincronizacoes_cadastros
            WHERE fonte = 'balanca'
            """
        )
        return {
            "fazendas_balanca": int(fazendas["total"] if fazendas else 0),
            "talhoes_balanca": int(talhoes["total"] if talhoes else 0),
            "sincronizado_em": sync["sincronizado_em"] if sync else None,
        }

    def close(self) -> None:
        with self._lock:
            conn = getattr(self, "conn", None)
            if conn is not None:
                conn.close()
                self.conn = None
