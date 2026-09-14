"""Reproducible fictional scenarios. No source database or file is read.

The seed is additive and runs once per database. Dates are anchored to setup day.
It never deletes records entered while exploring the demonstration.
"""
from __future__ import annotations

import csv
from datetime import date, datetime, timedelta, timezone
import hashlib
import json
from pathlib import Path
import random
import sys
import uuid

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import demo
import psycopg
from psycopg import sql
from agricola_shared.demo_safety import assert_demo_database_target

SETTINGS = demo.config()
TODAY = date.today()
RNG = random.Random(20260914)
FARMS = ["Fazenda Demo Aurora", "Fazenda Demo Horizonte", "Fazenda Demo Cedro", "Fazenda Demo Primavera",
         "Fazenda Demo Campo Azul", "Fazenda Demo Semente", "Fazenda Demo Boa Safra", "Fazenda Demo Planalto"]


def uid(value):
    return str(uuid.uuid5(uuid.NAMESPACE_URL, "https://example.invalid/operacoes-agricolas/" + value))


def insert(conn, table, **values):
    statement = sql.SQL("INSERT INTO {} ({}) VALUES ({}) ON CONFLICT DO NOTHING").format(
        sql.Identifier(table), sql.SQL(",").join(map(sql.Identifier, values)),
        sql.SQL(",").join(sql.Placeholder() for _ in values),
    )
    conn.execute(statement, list(values.values()))


def seed(module, callback):
    target = demo.database_url(SETTINGS, module)
    assert_demo_database_target(target)
    with psycopg.connect(target) as conn:
        conn.execute("CREATE TABLE IF NOT EXISTS oa_demo_seed (version text PRIMARY KEY, created_at date NOT NULL)")
        if conn.execute("SELECT 1 FROM oa_demo_seed WHERE version='v1'").fetchone():
            print(module + ": dados ficticios ja preparados")
            return
        callback(conn)
        insert(conn, "oa_demo_seed", version="v1", created_at=TODAY)
        print(module + ": cenario ficticio preparado")


def balance(conn):
    for i, name in enumerate(FARMS, 1):
        farm_id = uid(f"farm-{i}")
        code = f"{100 if i <= 4 else 200}-{900+i}"
        insert(conn, "farms", id=farm_id, code=code, name=name.upper(), section_name=name.upper(),
               owner_name="ORGANIZACAO FICTICIA DEMO", municipality="MUNICIPIO DEMONSTRATIVO",
               crop_year=str(TODAY.year), area_ha=80+i*8, area_alq=round((80+i*8)/2.42,4))
        for j in range(1, 5):
            insert(conn, "fields", id=uid(f"field-{i}-{j}"), code=str(j), name=f"Talhao demonstrativo {j}",
                   farm_id=farm_id, area_ha=20+i*2, area_alq=round((20+i*2)/2.42,4),
                   planted_area_ha=20+i*2, crop_year=str(TODAY.year), area_type="COLHEITA", active=True)
        order_id = uid(f"order-{i}")
        start = TODAY-timedelta(days=45)
        closed = i > 6
        insert(conn, "harvest_orders", id=order_id, number=str(90000+i), front_number=i if i<=4 else None,
               farm_id=farm_id, status="CLOSED" if closed else "ACTIVE", start_date=start,
               end_date=datetime.combine(TODAY-timedelta(days=1), datetime.min.time(), timezone.utc) if closed else None,
               origin="DEMO_FICTICIA")
        for j in range(1, 5):
            insert(conn, "harvest_order_fields", id=uid(f"order-field-{i}-{j}"), order_id=order_id, field_id=uid(f"field-{i}-{j}"))
        if i <= 4:
            insert(conn, "harvest_order_fronts", id=uid(f"front-{i}"), order_id=order_id, front_number=i)
        insert(conn, "harvest_order_history", id=uid(f"history-{i}"), order_id=order_id, order_number=str(90000+i),
               event_type="CLOSED" if closed else "CREATED", front_numbers_before="[]", front_numbers_after=json.dumps([i] if i<=4 else []))
    for day in range(28):
        reference = TODAY-timedelta(days=day)
        batch_id = uid(f"batch-{day}")
        rows = []
        for ticket in range(24):
            i = ticket % 8 + 1
            field = (ticket // 8 + day) % 4 + 1
            weight = round(RNG.uniform(23, 38), 3)
            status = "FIELD_NOT_RELEASED" if ticket == 22 else "OS_NOT_FOUND" if ticket == 23 else "OK"
            rows.append(dict(id=uid(f"ticket-{day}-{ticket}"), batch_id=batch_id, ticket_number=f"DEMO-{day:02d}-{ticket:03d}",
                             entry_date=reference, farm_id=uid(f"farm-{i}"), farm_name_raw=FARMS[i-1],
                             farm_code_raw=f"{100 if i<=4 else 200}-{900+i}", field_id=uid(f"field-{i}-{field}"),
                             field_code_raw="99" if status=="FIELD_NOT_RELEASED" else str(field),
                             order_id=None if status=="OS_NOT_FOUND" else uid(f"order-{i}"),
                             order_number_raw="99999" if status=="OS_NOT_FOUND" else str(90000+i),
                             vehicle_plate=f"DEMO-{ticket:03d}", gross_weight=weight+12, net_weight=weight,
                             status=status, notes="Registro inteiramente ficticio para demonstracao."))
        insert(conn, "import_batches", id=batch_id, file_name=f"pesagem-demo-{reference.isoformat()}.csv", source_type="CSV",
               report_date=reference, period_start=reference, period_end=reference, total_net_weight=sum(r["net_weight"] for r in rows),
               total_trips=len(rows), row_count=len(rows), ok_count=22, error_count=2)
        for row in rows: insert(conn, "cane_entries", **row)
    for i in range(1, 5):
        insert(conn, "fleet_movements", id=uid(f"movement-{i}"), movement_type="LEFT_MILL",
               equipment_code=str(9000+i), equipment_type="COLHEDORA", front_number=i,
               reason="Manutencao preventiva simulada", status="OPEN")


def people(conn):
    for i in range(1, 33):
        function = ["MOTORISTA", "OPERADOR DE COLHEDORA", "LIDER DE FRENTE", "OPERADOR DE TRANSBORDO"][(i-1)%4]
        expiration = TODAY + timedelta(days=365 if i%5 else -15 if i%10 else 20)
        insert(conn, "colaboradores", codigo_colaborador=str(90000+i), codigo_interno=f"DEMO-{i:03d}",
               nome=f"PESSOA DEMONSTRATIVA {i:02d}", situacao="ATIVO", modalidade="SAFRISTA",
               cidade="MUNICIPIO DEMONSTRATIVO", municipio="MUNICIPIO DEMONSTRATIVO",
               validade_cnh=expiration.isoformat(), categoria_cnh="E" if i%4==1 else "D",
               frente_safra=str((i-1)%4+1), funcao_safra=function, funcao=function, turno_safra="A" if i%2 else "B",
               horario="07:00 - 15:00" if i%2 else "15:00 - 23:00", data_admissao=(TODAY-timedelta(days=120)).isoformat(),
               tem_foto="NÃO", observacao_1="Pessoa inteiramente ficticia. Sem CPF, RG, telefone ou documento real.",
               gestor_responsavel="GESTOR DEMONSTRATIVO")


def notes(conn):
    for i in range(1, 33):
        insert(conn, "motoristas", codigo=90000+i, nome=f"PESSOA DEMONSTRATIVA {i:02d}")
    for i, name in enumerate(FARMS,1):
        insert(conn, "fazendas", codigo=f"{100 if i<=4 else 200}-{900+i}", nome=name.upper())
    for i in range(1,4): insert(conn, "variedades", id=i, nome=f"VARIEDADE DEMO {i}")
    for i in range(1,85):
        farm=(i-1)%4+1
        driver=(i-1)%8*4+1
        operator=(i-1)%8*4+2
        harvested=TODAY-timedelta(days=i%21)
        planted=harvested+timedelta(days=1 if i%5 else 2)
        insert(conn, "notas", numero=900000+i, motorista_cod=90000+driver, motorista_nome=f"PESSOA DEMONSTRATIVA {driver:02d}",
               operador_cod=90000+operator, operador_nome=f"PESSOA DEMONSTRATIVA {operator:02d}", caminhao=str(9800+i%8),
               colhedora=str(9000+farm), faz_muda_cod=f"100-{900+farm}", faz_muda_nome=FARMS[farm-1].upper(),
               talhao=str(i%4+1), faz_plantio_cod=f"200-{904+farm}", faz_plantio_nome=FARMS[farm+3].upper(),
               variedade_id=i%3+1, variedade_nome=f"VARIEDADE DEMO {i%3+1}",
               data_colheita=harvested.isoformat(), data_plantio=planted.isoformat(), duplicado=0)


def analytics(conn):
    reasons = ["Abastecimento", "Manutencao preventiva", "Troca de turno", "Aguardando transporte", "Ajuste operacional"]
    for front in range(1,5):
        for equipment in range(3):
            kind = "TRANSBORDO" if equipment == 2 else "COLHEDORA"
            insert(conn, "frotas_por_frente", frente=str(front), frota=f"{kind} {9000+front*10+equipment}")
    for day in range(28):
        reference=(TODAY-timedelta(days=day)).strftime("%d-%m-%Y")
        for front in range(1,5):
            minutes=15+(front*7+day*3)%70
            insert(conn, "relatorio_operacao_diaria", data=reference, frente=str(front), turno="A" if day%2 else "B",
                   frota=f"{'TRANSBORDO' if day%3 == 2 else 'COLHEDORA'} {9000+front*10+day%3}", motivo=reasons[(day+front)%len(reasons)], parou_hora="09:00",
                   voltou_hora=f"{9+minutes//60:02d}:{minutes%60:02d}", total_hora_parado=f"{minutes//60:02d}:{minutes%60:02d}",
                   eficiencia=round(100*(600-minutes)/600,2), fundo_agricola=FARMS[front-1].upper(),
                   chuva="Não", incendio="Não", status_parada="Finalizada")
            insert(conn, "colheita_mecanizada", data=reference, frente=str(front), turno="A",
                   fazenda=FARMS[front-1].upper(), area_colhida=round(8+RNG.random()*5,2),
                   produtividade=round(70+RNG.random()*20,2), viagens=16+day%8)


def fixture():
    # A small import file generated only from explicit synthetic constants.
    target = ROOT / "demo-data"
    target.mkdir(exist_ok=True)
    csv_path = target / "pesagem-ficticia.csv"
    if not csv_path.exists():
        with csv_path.open("w", encoding="utf-8", newline="") as handle:
            writer=csv.writer(handle,delimiter=";")
            writer.writerow(["ticket", "data", "fazenda", "talhao", "os", "placa", "peso_bruto", "peso_liquido"])
            for i in range(1,7):
                writer.writerow([f"EXEMPLO-{i}",TODAY.strftime("%d/%m/%Y"),"100-901",str(i%4+1),"90001",f"DEMO-{i:03d}","43,500","31,500"])
    # Fleet workbook is a technical input fixture for the retained Excel parser.
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
    fleet_file = demo.STATE / "frotas-ficticias.xlsx"
    if not fleet_file.exists():
        wb=Workbook(); ws=wb.active; ws.title="Frota ficticia"
        ws.append(["COLHEDORA", "TRANSBORDO", "CARREGADEIRA", "VIVENCIA", "CAMINHAO D'AGUA", "FURGAO"])
        for front in range(1,5):
            ws.append([f"FRENTE {front}"])
            for item in range(3):
                ws.append([str(9000+front*10+item),str(9500+front*10+item),
                           str(9600+front) if item == 0 else "", str(9700+front) if item == 0 else "",
                           str(9800+front) if item == 0 else "", str(9900+front) if item == 0 else ""])
                for cell in ws[ws.max_row]:
                    if cell.value:
                        cell.fill = PatternFill("solid", fgColor="FFFF0000" if front == 4 and item == 2 else "FF92D050")
        wb.save(fleet_file)


if __name__ == "__main__":
    seed("balanca", balance)
    seed("colaboradores", people)
    seed("notas", notes)
    seed("analises", analytics)
    fixture()
